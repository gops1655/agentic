from email.message import EmailMessage
from pathlib import Path

from openpyxl import load_workbook

from invoice_agent import agent as agent_mod
from invoice_agent.agent import InvoiceAgent
from invoice_agent.config import Settings
from invoice_agent.excel_writer import SourceInvoice, reconcile, save_workbook
from invoice_agent.extractor import ExtractionError
from invoice_agent.mailer import build_reply, matches_filters, parse_email
from invoice_agent.models import ExtractionResult, Invoice, LineItem, TaxLine
from invoice_agent.store import Store


def make_invoice(**overrides) -> Invoice:
    base = dict(
        document_type="tax invoice", invoice_number="INV-001", invoice_date="2026-09-01", due_date=None,
        po_number="PO-9", vendor_name="Acme Medical Supplies", vendor_address="Hyderabad",
        vendor_tax_id="36ABCDE1234F1Z5", vendor_pan=None, bill_to_name="Yashoda Hospitals",
        bill_to_address=None, bill_to_tax_id=None, place_of_supply="Telangana", currency="INR",
        subtotal=300.0, discount_total=None, tax_total=54.0, shipping=None, round_off=None,
        total_amount=354.0, amount_paid=None, balance_due=None, payment_terms="30 days",
        bank_details=None, notes=None,
        line_items=[
            LineItem(line_no=1, item_code=None, description="Gloves", hsn_sac="4015", batch_no="B1",
                     expiry_date="2028-01-01", quantity=10, unit="Box", unit_price=20, discount=None,
                     tax_rate_percent=18, tax_amount=36, amount=200),
            LineItem(line_no=2, item_code=None, description="Masks", hsn_sac="6307", batch_no=None,
                     expiry_date=None, quantity=5, unit="Box", unit_price=20, discount=None,
                     tax_rate_percent=18, tax_amount=18, amount=100),
        ],
        taxes=[TaxLine(name="CGST", rate_percent=9, taxable_amount=300, amount=27),
               TaxLine(name="SGST", rate_percent=9, taxable_amount=300, amount=27)],
    )
    base.update(overrides)
    return Invoice(**base)


def make_raw_email(sender="Vendor <billing@acme.in>", subject="Invoice INV-001") -> bytes:
    msg = EmailMessage()
    msg["From"] = sender
    msg["To"] = "ap@hospital.com"
    msg["Subject"] = subject
    msg["Message-ID"] = "<abc@acme.in>"
    msg.set_content("Please find attached our invoice.")
    msg.add_attachment(b"%PDF-1.4 fake", maintype="application", subtype="pdf", filename="INV-001.pdf")
    return msg.as_bytes()


def settings(**kw) -> Settings:
    return Settings(email_address="ap@hospital.com", email_password="x", anthropic_api_key="k", **kw)


def test_reconcile():
    assert reconcile(make_invoice())[1] == "OK"
    assert reconcile(make_invoice(subtotal=999, total_amount=1200))[1].startswith("Mismatch")


def test_workbook(tmp_path: Path):
    path = save_workbook([SourceInvoice("a.pdf", make_invoice())], tmp_path / "out.xlsx", ["note"])
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Line Items", "Taxes", "Extraction Notes"]
    summary = wb["Summary"]
    assert summary["C2"].value == "INV-001"
    assert summary.cell(row=3, column=1).value == "TOTAL"
    lines = wb["Line Items"]
    assert lines.max_row == 4  # header + 2 lines + total
    assert lines["E2"].value == "Gloves"


def test_parse_and_reply():
    mail = parse_email("7", make_raw_email())
    assert mail.sender_email == "billing@acme.in"
    assert len(mail.pdfs) == 1 and mail.pdfs[0].filename == "INV-001.pdf"
    reply = build_reply(settings(), mail, "hello", [])
    assert reply["To"] == "billing@acme.in"
    assert reply["Subject"] == "Re: Invoice INV-001"
    assert reply["In-Reply-To"] == "<abc@acme.in>"


def test_filters():
    mail = parse_email("1", make_raw_email())
    assert matches_filters(mail, settings())
    assert matches_filters(mail, settings(allowed_senders=["acme.in"]))
    assert not matches_filters(mail, settings(allowed_senders=["other.com"]))
    assert not matches_filters(mail, settings(subject_keywords=["bill"]))
    own = parse_email("2", make_raw_email(sender="ap@hospital.com"))
    assert not matches_filters(own, settings())


class FakeExtractor:
    def __init__(self, error=None):
        self.error = error

    def extract(self, data, filename):
        if self.error:
            raise self.error
        return ExtractionResult(invoices=[make_invoice()], warnings=[])


def make_agent(tmp_path, monkeypatch, extractor):
    monkeypatch.setattr(agent_mod, "OUTPUT_DIR", tmp_path)
    a = InvoiceAgent(settings(), store=Store(tmp_path / "s.db"), log=lambda *_: None)
    a.extractor = extractor
    return a


def test_convert_email(tmp_path, monkeypatch):
    a = make_agent(tmp_path, monkeypatch, FakeExtractor())
    mail = parse_email("7", make_raw_email())
    result = a.convert_email(mail)
    assert result.ok and result.excel_path.exists()
    assert "INV-001" in a.reply_body(mail, result)
    a.record(mail, result, "sent")
    assert a.store.seen("<abc@acme.in>")


def test_retryable_error_does_not_fail_file(tmp_path, monkeypatch):
    a = make_agent(tmp_path, monkeypatch, FakeExtractor(ExtractionError("rate limit", retryable=True)))
    result = a.convert_email(parse_email("7", make_raw_email()))
    assert result.retry_later and not result.failed_files and result.excel_path is None
