"""Turn extracted invoices into a formatted Excel workbook."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from .models import Invoice

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TOTAL_FONT = Font(bold=True)
WARN_FILL = PatternFill("solid", fgColor="FCE4D6")
MONEY = "#,##0.00"
QTY = "#,##0.###"


@dataclass
class SourceInvoice:
    """An extracted invoice plus the PDF it came from."""

    source_file: str
    invoice: Invoice


SUMMARY_COLUMNS = [
    ("Source PDF", "source", None, 28),
    ("Document Type", "document_type", None, 14),
    ("Invoice No", "invoice_number", None, 18),
    ("Invoice Date", "invoice_date", None, 13),
    ("Due Date", "due_date", None, 13),
    ("PO Number", "po_number", None, 16),
    ("Vendor", "vendor_name", None, 30),
    ("Vendor GSTIN / Tax ID", "vendor_tax_id", None, 20),
    ("Vendor PAN", "vendor_pan", None, 13),
    ("Vendor Address", "vendor_address", None, 40),
    ("Bill To", "bill_to_name", None, 28),
    ("Bill To GSTIN / Tax ID", "bill_to_tax_id", None, 20),
    ("Place of Supply", "place_of_supply", None, 16),
    ("Currency", "currency", None, 9),
    ("Subtotal", "subtotal", MONEY, 14),
    ("Discount", "discount_total", MONEY, 12),
    ("Tax Total", "tax_total", MONEY, 13),
    ("Shipping / Other", "shipping", MONEY, 13),
    ("Round Off", "round_off", MONEY, 10),
    ("Total Amount", "total_amount", MONEY, 15),
    ("Amount Paid", "amount_paid", MONEY, 13),
    ("Balance Due", "balance_due", MONEY, 14),
    ("Payment Terms", "payment_terms", None, 18),
    ("Bank Details", "bank_details", None, 34),
    ("Line Items", "line_count", None, 10),
    ("Line Items Sum", "line_sum", MONEY, 15),
    ("Check", "check", None, 34),
    ("Notes", "notes", None, 40),
]

LINE_COLUMNS = [
    ("Invoice No", 18, None),
    ("Vendor", 28, None),
    ("Line", 6, None),
    ("Item Code", 14, None),
    ("Description", 48, None),
    ("HSN/SAC", 11, None),
    ("Batch", 12, None),
    ("Expiry", 11, None),
    ("Qty", 9, QTY),
    ("Unit", 8, None),
    ("Unit Price", 13, MONEY),
    ("Discount", 11, MONEY),
    ("Tax %", 7, "0.##"),
    ("Tax Amount", 13, MONEY),
    ("Amount", 14, MONEY),
]

TAX_COLUMNS = [
    ("Invoice No", 18, None),
    ("Vendor", 28, None),
    ("Tax", 12, None),
    ("Rate %", 8, "0.##"),
    ("Taxable Amount", 16, MONEY),
    ("Tax Amount", 14, MONEY),
]


def reconcile(inv: Invoice) -> tuple[float | None, str]:
    """Sum the line items and compare them with the invoice's own totals."""
    amounts = [li.amount for li in inv.line_items if li.amount is not None]
    if not amounts:
        return None, "No line amounts" if inv.line_items else "No line items"
    line_sum = round(sum(amounts), 2)
    targets = [t for t in (inv.subtotal, inv.total_amount) if t is not None]
    if not targets:
        return line_sum, "No totals printed"
    tolerance = max(1.0, 0.005 * max(abs(t) for t in targets))
    if any(abs(line_sum - t) <= tolerance for t in targets):
        return line_sum, "OK"
    return line_sum, f"Mismatch: lines {line_sum:,.2f} vs subtotal/total"


def _header(ws: Worksheet, titles: list[str], widths: list[int]) -> None:
    ws.append(titles)
    for idx, width in enumerate(widths, start=1):
        cell = ws.cell(row=1, column=idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"


def _format_column(ws: Worksheet, col: int, fmt: str | None, first: int, last: int) -> None:
    if not fmt:
        return
    for row in range(first, last + 1):
        ws.cell(row=row, column=col).number_format = fmt


def _totals_row(ws: Worksheet, label_col: int, sum_cols: list[int], data_rows: int) -> None:
    if data_rows == 0:
        return
    row = data_rows + 2
    ws.cell(row=row, column=label_col, value="TOTAL").font = TOTAL_FONT
    for col in sum_cols:
        letter = get_column_letter(col)
        cell = ws.cell(row=row, column=col, value=f"=SUBTOTAL(9,{letter}2:{letter}{row - 1})")
        cell.font = TOTAL_FONT
        cell.number_format = MONEY


def build_workbook(items: list[SourceInvoice], warnings: list[str] | None = None) -> Workbook:
    wb = Workbook()

    # --- Summary: one row per invoice ---
    ws = wb.active
    ws.title = "Summary"
    _header(ws, [c[0] for c in SUMMARY_COLUMNS], [c[3] for c in SUMMARY_COLUMNS])
    for item in items:
        inv = item.invoice
        line_sum, check = reconcile(inv)
        values = inv.model_dump()
        values.update(source=item.source_file, line_count=len(inv.line_items), line_sum=line_sum, check=check)
        ws.append([values.get(key) for _, key, _, _ in SUMMARY_COLUMNS])
        if check.startswith("Mismatch"):
            ws.cell(row=ws.max_row, column=len(SUMMARY_COLUMNS) - 1).fill = WARN_FILL
    for idx, (_, _, fmt, _) in enumerate(SUMMARY_COLUMNS, start=1):
        _format_column(ws, idx, fmt, 2, len(items) + 1)
    money_cols = [i for i, c in enumerate(SUMMARY_COLUMNS, start=1) if c[2] == MONEY and c[1] != "line_sum"]
    _totals_row(ws, 1, money_cols, len(items))
    if items:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(SUMMARY_COLUMNS))}{len(items) + 1}"

    # --- Line items: every row from every invoice ---
    wl = wb.create_sheet("Line Items")
    _header(wl, [c[0] for c in LINE_COLUMNS], [c[1] for c in LINE_COLUMNS])
    n_lines = 0
    for item in items:
        inv = item.invoice
        for li in inv.line_items:
            wl.append([
                inv.invoice_number, inv.vendor_name, li.line_no, li.item_code, li.description,
                li.hsn_sac, li.batch_no, li.expiry_date, li.quantity, li.unit, li.unit_price,
                li.discount, li.tax_rate_percent, li.tax_amount, li.amount,
            ])
            wl.cell(row=wl.max_row, column=5).alignment = Alignment(wrap_text=True, vertical="top")
            n_lines += 1
    for idx, (_, _, fmt) in enumerate(LINE_COLUMNS, start=1):
        _format_column(wl, idx, fmt, 2, n_lines + 1)
    _totals_row(wl, 1, [14, 15], n_lines)
    if n_lines:
        wl.auto_filter.ref = f"A1:{get_column_letter(len(LINE_COLUMNS))}{n_lines + 1}"

    # --- Taxes breakdown ---
    wt = wb.create_sheet("Taxes")
    _header(wt, [c[0] for c in TAX_COLUMNS], [c[1] for c in TAX_COLUMNS])
    n_tax = 0
    for item in items:
        inv = item.invoice
        for t in inv.taxes:
            wt.append([inv.invoice_number, inv.vendor_name, t.name, t.rate_percent, t.taxable_amount, t.amount])
            n_tax += 1
    for idx, (_, _, fmt) in enumerate(TAX_COLUMNS, start=1):
        _format_column(wt, idx, fmt, 2, n_tax + 1)
    _totals_row(wt, 1, [6], n_tax)

    # --- Notes from extraction ---
    if warnings:
        wn = wb.create_sheet("Extraction Notes")
        _header(wn, ["Note"], [110])
        for w in warnings:
            wn.append([w])
            wn.cell(row=wn.max_row, column=1).alignment = Alignment(wrap_text=True)

    return wb


def save_workbook(items: list[SourceInvoice], path: Path, warnings: list[str] | None = None) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    build_workbook(items, warnings).save(path)
    return path
